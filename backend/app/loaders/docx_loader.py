"""loaders.docx_loader
Word document loader backed by python-docx. Extracts paragraph text and
table cell text so indicators inside tables are not lost.
"""

import docx

from app.loaders.base import BaseLoader


class DocxLoader(BaseLoader):
    format_name = "docx"
    extensions = (".docx",)

    def load(self, filepath: str) -> str:
        document = docx.Document(filepath)
        parts = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text)
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
